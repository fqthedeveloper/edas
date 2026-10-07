import os
import re
import io
import docx
import pypdf
from typing import Dict, List, Any, Tuple


def slugify_key(text: str) -> str:
    """Convert label text to clean snake_case identifier."""
    text = text.strip().lower()
    text = re.sub(r'[^\w\s-]', '', text)
    text = re.sub(r'[-\s]+', '_', text)
    return text.strip('_') or 'field'


def infer_field_type(label: str, sample_val: str = '') -> str:
    """Infer the dynamic field type based on label name and sample value."""
    lbl = label.lower()
    val = sample_val.lower().strip()

    if any(k in lbl for k in ['date', 'dated', 'validity']):
        return 'date'
    if any(k in lbl for k in ['rate', 'flow', 'velocity', 'pressure', 'temp', 'drop', 'efficiency', 'capacity', 'qty', 'quantity', 'amount']):
        return 'number'
    if any(k in lbl for k in ['type', 'class', 'equipment', 'status', 'moc']):
        return 'dropdown'
    if '\n' in val or len(val) > 100:
        return 'textarea'
    return 'text'


class DocumentLayoutAnalyzer:
    """
    Analyzes Word (.docx) and PDF files to extract:
    1. Metadata key-value fields
    2. Structured sections
    3. Tables and tabular columns
    4. Auto-generated HTML/CSS template code matching the layout
    5. TemplateField model definitions
    """

    def __init__(self, file_path_or_bytes, filename: str):
        self.filename = filename
        self.is_docx = filename.lower().endswith('.docx')
        self.is_pdf = filename.lower().endswith('.pdf')
        self.raw_data = file_path_or_bytes

    def parse(self) -> Dict[str, Any]:
        if self.is_docx:
            return self._parse_docx()
        elif self.is_pdf:
            return self._parse_pdf()
        else:
            raise ValueError("Unsupported format. Please upload a .docx or .pdf file.")

    def _parse_docx(self) -> Dict[str, Any]:
        if isinstance(self.raw_data, (str, os.PathLike)):
            doc = docx.Document(self.raw_data)
        else:
            doc = docx.Document(io.BytesIO(self.raw_data))

        title = "Document Template"
        fields = []
        seen_keys = set()
        sections = []

        # 1. Parse paragraphs for Title and Key-Values
        metadata_pairs = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            # First large heading or standalone uppercase paragraph
            if not metadata_pairs and len(text) < 40 and text.isupper():
                title = text
                continue

            # Look for {{ placeholder }} patterns
            curly_matches = re.findall(r'\{\{\s*(\w+)\s*\}\}', text)
            for m in curly_matches:
                k = slugify_key(m)
                if k not in seen_keys:
                    seen_keys.add(k)
                    fields.append({
                        'field_name': k,
                        'label': m.replace('_', ' ').title(),
                        'field_type': infer_field_type(m),
                        'section': 'General Information',
                        'sample_value': ''
                    })

            # Look for Label: Value or Label: {{ var }}
            colon_splits = re.split(r'[:\-]\s+', text)
            if len(colon_splits) >= 2:
                for match in re.finditer(r'([A-Za-z0-9\s/]+)\s*[:\-]\s*([^\n\r]+?)(?=[A-Z\s]{4,}[:\-]|$)', text):
                    lbl = match.group(1).strip()
                    val = match.group(2).strip()
                    if 2 <= len(lbl) <= 40 and not lbl.startswith('http'):
                        k = slugify_key(lbl)
                        if k not in seen_keys:
                            seen_keys.add(k)
                            fields.append({
                                'field_name': k,
                                'label': lbl.title(),
                                'field_type': infer_field_type(lbl, val),
                                'section': 'General Information',
                                'sample_value': val
                            })

        # 2. Parse Tables
        table_html_rows = []
        current_section = 'Product Details'

        for t_idx, table in enumerate(doc.tables):
            # Process table row by row
            for row in table.rows:
                # Deduplicate merged cells
                unique_cells = []
                seen_cell_text = set()
                for c in row.cells:
                    ctext = c.text.strip()
                    if c._tc not in [uc._tc for uc in unique_cells]:
                        unique_cells.append(c)

                if not unique_cells:
                    continue

                # Section banner check
                first_text = unique_cells[0].text.strip()
                if len(unique_cells) == 1 or all(c.text.strip() == first_text for c in unique_cells):
                    if first_text.isupper() and len(first_text) > 3:
                        current_section = first_text.title()
                        if current_section not in sections:
                            sections.append(current_section)
                        table_html_rows.append(
                            f'<tr><td colspan="4" class="section-title">{first_text}</td></tr>'
                        )
                        continue

                # Header vs Data Row
                is_header_row = any(c.text.strip().isupper() and len(c.text.strip()) > 2 for c in unique_cells)
                
                # Check for label/value pattern inside table cells
                row_cells_html = []
                col_width = int(100 / max(len(unique_cells), 1))

                for cell in unique_cells:
                    cell_text = cell.text.strip()
                    # Check for curly brace in cell
                    curly = re.search(r'\{\{\s*(\w+)\s*\}\}', cell_text)
                    if curly:
                        k = slugify_key(curly.group(1))
                        if k not in seen_keys:
                            seen_keys.add(k)
                            fields.append({
                                'field_name': k,
                                'label': curly.group(1).replace('_', ' ').title(),
                                'field_type': infer_field_type(k),
                                'section': current_section,
                                'sample_value': ''
                            })
                        row_cells_html.append(f'<td>{{{{ {k}|default:"-" }}}}</td>')
                    elif is_header_row:
                        lbl = cell_text.strip()
                        k = slugify_key(lbl)
                        row_cells_html.append(f'<th width="{col_width}%">{lbl}</th>')
                    else:
                        row_cells_html.append(f'<td>{cell_text or "-"}</td>')

                if is_header_row:
                    table_html_rows.append(f'<tr>{"".join(row_cells_html)}</tr>')
                else:
                    table_html_rows.append(f'<tr>{"".join(row_cells_html)}</tr>')

        # 3. Generate HTML/CSS
        html_layout = self._build_html_template(title, fields, table_html_rows)

        return {
            'template_name': title.title(),
            'fields': fields,
            'sections': sections or ['General Information', 'Product Details', 'Test Data'],
            'html_layout': html_layout
        }

    def _parse_pdf(self) -> Dict[str, Any]:
        if isinstance(self.raw_data, (str, os.PathLike)):
            reader = pypdf.PdfReader(self.raw_data)
        else:
            reader = pypdf.PdfReader(io.BytesIO(self.raw_data))

        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"

        title = "Imported Document"
        lines = [l.strip() for l in text.splitlines() if l.strip()]

        if lines:
            for l in lines[:5]:
                if len(l) < 40 and ('CERTIFICATE' in l or 'REPORT' in l or 'INVOICE' in l):
                    title = l
                    break

        fields = []
        seen_keys = set()
        
        # Regex for common label: value patterns
        patterns = [
            r'([A-Z\s]{3,25})\s*:\s*([^\n\r]+?)(?=[A-Z\s]{4,}\s*:|$)',
            r'([A-Za-z\s]{3,25})\s*[:\-]\s*([^\n\r]+)'
        ]
        
        for pat in patterns:
            for m in re.finditer(pat, text):
                lbl = m.group(1).strip()
                val = m.group(2).strip()
                if 2 <= len(lbl) <= 30 and not lbl.startswith('http'):
                    k = slugify_key(lbl)
                    if k not in seen_keys:
                        seen_keys.add(k)
                        fields.append({
                            'field_name': k,
                            'label': lbl.title(),
                            'field_type': infer_field_type(lbl, val),
                            'section': 'General Information',
                            'sample_value': val[:50]
                        })

        html_layout = self._build_html_template(title, fields, [])
        return {
            'template_name': title.title(),
            'fields': fields,
            'sections': ['General Information', 'Specifications', 'Results'],
            'html_layout': html_layout
        }

    def _build_html_template(self, title: str, fields: List[Dict], table_rows: List[str]) -> str:
        """Synthesize responsive A4-accurate print layout."""
        # Top metadata fields
        meta_html = ""
        paired_fields = [f for f in fields if f['section'] == 'General Information']
        
        for i in range(0, len(paired_fields), 2):
            f1 = paired_fields[i]
            f2 = paired_fields[i+1] if i+1 < len(paired_fields) else None
            meta_html += '<tr>\n'
            meta_html += f'    <td class="left"><span class="label">{f1["label"]}:</span> {{{{ {f1["field_name"]}|default:"-" }}}}</td>\n'
            if f2:
                meta_html += f'    <td class="right"><span class="label">{f2["label"]}:</span> {{{{ {f2["field_name"]}|default:"-" }}}}</td>\n'
            else:
                meta_html += '    <td></td>\n'
            meta_html += '</tr>\n'

        tables_block = "\n".join(table_rows) if table_rows else """
            <tr><td colspan="4" class="section-title">SPECIFICATIONS</td></tr>
            <tr>
                <th width="25%">PARAMETER</th>
                <th width="25%">STANDARD</th>
                <th width="25%">OBSERVED</th>
                <th width="25%">REMARKS</th>
            </tr>
            <tr>
                <td>Sample Parameter</td>
                <td>{{ standard|default:"-" }}</td>
                <td>{{ observed|default:"-" }}</td>
                <td>{{ remarks|default:"Pass" }}</td>
            </tr>
        """

        return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8">
<title>{title}</title>
<style>
@page {{
    size: A4 portrait;
    margin: 10mm 14mm 8mm 14mm;
}}
html, body {{
    margin: 0;
    padding: 0;
}}
body {{
    font-family: Arial, Helvetica, sans-serif;
    font-size: 10pt;
    color: #000;
    line-height: 1.35;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
}}
.watermark {{
    position: fixed;
    top: 50%;
    left: 50%;
    transform: translate(-50%, -50%);
    width: 65%;
    text-align: center;
    z-index: 0;
    pointer-events: none;
    opacity: 0.25;
}}
.watermark img {{
    max-width: 100%;
    height: auto;
}}
.page {{
    position: relative;
    z-index: 2;
    box-sizing: border-box;
}}
.header-table {{
    width: 100%;
    border-collapse: collapse;
    margin-bottom: 4px;
}}
.header-table td {{
    border: none;
    padding: 0;
    vertical-align: middle;
}}
.logo-cell {{
    width: 35%;
    text-align: left;
}}
.logo-cell img {{
    height: 52px;
    width: auto;
}}
.brand-cell {{
    width: 65%;
    text-align: right;
    font-size: 22pt;
    font-weight: bold;
    color: #0c7fb8;
    letter-spacing: 1.5px;
}}
.cert-title {{
    text-align: center;
    font-size: 17pt;
    font-weight: bold;
    color: #0c7fb8;
    margin: 6px 0 10px;
    letter-spacing: 1px;
    border-bottom: 2px solid #0c7fb8;
    padding-bottom: 4px;
}}
.info-table {{
    width: 100%;
    border-collapse: collapse;
    margin-bottom: 12px;
}}
.info-table td {{
    border: none;
    padding: 3px 0;
    font-size: 10pt;
    vertical-align: top;
}}
.left {{ text-align: left; }}
.right {{ text-align: right; }}
.label {{ font-weight: bold; color: #222; }}
.main-table {{
    width: 100%;
    border-collapse: collapse;
    font-size: 9.5pt;
    table-layout: fixed;
}}
.main-table th, .main-table td {{
    border: 1px solid #000;
    padding: 7px 6px;
    text-align: center;
    vertical-align: middle;
    box-sizing: border-box;
    word-wrap: break-word;
}}
.main-table th {{
    font-weight: bold;
    font-size: 9pt;
    letter-spacing: 0.3px;
}}
.section-title {{
    font-size: 10.5pt;
    font-weight: bold;
    text-transform: uppercase;
    color: #0c7fb8;
    text-align: center !important;
    padding: 6px !important;
    background: #f3fafd;
    letter-spacing: 0.5px;
}}
.closing-container {{
    margin-top: 14px;
    width: 100%;
    border: 1px solid #000;
    box-sizing: border-box;
    padding: 10px 12px;
    background: #fff;
}}
.closing-table {{
    width: 100%;
    border: none;
    border-collapse: collapse;
    margin: 0;
}}
.closing-table td {{
    border: none;
    padding: 0;
    vertical-align: bottom;
}}
.contact-line {{
    margin-top: 10px;
    font-size: 9.5pt;
    text-align: left;
    border-top: 1px dashed #bbb;
    padding-top: 8px;
}}
.footer {{
    margin-top: 14px;
    padding-top: 8px;
    border-top: 2px solid #0c7fb8;
    text-align: center;
    font-size: 9.5pt;
    color: #000;
}}
.footer a {{
    color: #000;
    text-decoration: none;
    font-weight: bold;
}}
</style>
</head>
<body>
{{% if has_watermark and watermark_b64 %}}
<div class="watermark">
    <img src="{{{{ watermark_b64 }}}}">
</div>
{{% endif %}}

<div class="page">
    <table class="header-table">
        <tr>
            <td class="logo-cell">
                {{% if logo_b64 %}}
                    <img src="{{{{ logo_b64 }}}}">
                {{% elif company.logo %}}
                    <img src="{{{{ base_url }}}}{{{{ company.logo.url }}}}">
                {{% endif %}}
            </td>
            <td class="brand-cell">
                {{{{ company.name|default:"ENTERPRISE DOCUMENT" }}}}
            </td>
        </tr>
    </table>

    <div class="cert-title">
        {title}
    </div>

    <table class="info-table">
        {meta_html}
    </table>

    <table class="main-table">
        {tables_block}
    </table>

    <div class="closing-container">
        <table class="closing-table">
            <tr>
                <td style="width:30%; text-align:left;">
                    {{% if qr_b64 %}}
                        <img src="{{{{ qr_b64 }}}}" style="height:68px; width:auto; display:block;">
                    {{% elif document.qr_code %}}
                        <img src="{{{{ base_url }}}}{{{{ document.qr_code.url }}}}" style="height:68px; width:auto; display:block;">
                    {{% endif %}}
                </td>
                <td style="width:35%;"></td>
                <td style="width:35%; text-align:center;">
                    <div style="font-weight:bold; margin-bottom:5px; font-size:10pt;">
                        INSPECTED BY
                    </div>
                    {{% if signature_b64 %}}
                        <img src="{{{{ signature_b64 }}}}" style="height:52px; width:auto; display:inline-block; mix-blend-mode: multiply;">
                    {{% elif company.signature %}}
                        <img src="{{{{ base_url }}}}{{{{ company.signature.url }}}}" style="height:52px; width:auto; display:inline-block; mix-blend-mode: multiply;">
                    {{% else %}}
                        <div style="height:52px;"></div>
                    {{% endif %}}
                    <div style="font-weight:bold; margin-top:3px; font-size:10pt;">
                        Authorised Signatory
                    </div>
                </td>
            </tr>
        </table>
        <div class="contact-line">
            For Queries Related To This Document Please Contact : <strong>{{{{ company.email|default:"support@company.com" }}}}</strong>
        </div>
    </div>

    <div class="footer">
        <div style="font-size:11pt; font-weight:bold; margin-bottom:2px; color:#0c7fb8;">
            {{{{ company.legal_name|default:company.name }}}}
        </div>
        <div style="font-size:9.5pt; color:#333;">
            Email : <a href="mailto:{{{{ company.email }}}}">{{{{ company.email }}}}</a>
            &nbsp;&nbsp;&nbsp;|&nbsp;&nbsp;&nbsp;
            Website : <a href="{{{{ company.website }}}}" target="_blank">{{{{ company.website }}}}</a>
        </div>
    </div>
</div>
</body>
</html>"""

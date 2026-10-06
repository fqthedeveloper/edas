import datetime
import json
import io
import zipfile
import logging
import openpyxl
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl.writer.excel import save_virtual_workbook
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.views.decorators.http import require_POST
from django.utils.dateparse import parse_date
from django.template import Template as DjangoTemplate, Context
from django.views.decorators.csrf import csrf_exempt
from django.conf import settings

from clients.models import Client
from products.models import Product
from templatesapp.models import Template, TemplateField
from .models import Document, DocumentValue
from .forms import build_dynamic_form, ExcelUploadForm
from sequences.services import get_next_number
from .tasks import generate_document_pdf_docx, _generate_document

logger = logging.getLogger(__name__)


# ---------- STEP 1 ----------
@login_required
def document_wizard_step1(request):
    clients = Client.objects.filter(is_active=True)
    return render(request, 'documents/step1.html', {'clients': clients})


# ---------- STEP 2 ----------
@login_required
def document_wizard_step2(request):
    client_id = request.GET.get('client')
    if not client_id:
        return redirect('document_step1')
    client = get_object_or_404(Client, id=client_id)
    products = Product.objects.filter(company=client.company, default_template__isnull=False)
    return render(request, 'documents/step2.html', {'client': client, 'products': products})


@login_required
def document_wizard_step3(request):
    client_id = request.GET.get('client')
    product_id = request.GET.get('product')
    if not client_id or not product_id:
        return redirect('document_step1')

    client = get_object_or_404(Client, id=client_id)
    product = get_object_or_404(Product, id=product_id)
    template = product.default_template or client.default_template

    if not template:
        return render(request, 'documents/error.html', {'msg': 'No template defined for this product/client.'})

    # Build initial data
    initial = {}
    initial['client_name'] = client.name
    initial['client_address'] = client.address
    initial['client_email'] = client.email
    initial['client_phone'] = client.phone
    initial['client_gst'] = client.gst
    initial['client_contact_person'] = client.contact_person

    if product.default_fields:
        for key, val in product.default_fields.items():
            initial[key] = val

    today = datetime.date.today()
    initial['date'] = today.isoformat()
    initial['testing_date'] = today.isoformat()

    # TC Number Reservation via Session
    prefix = 'AFPL/2026-27/'
    session_key = f'tc_reserved_{prefix}'
    if session_key not in request.session:
        next_num = get_next_number(prefix)
        request.session[session_key] = f"{prefix}{next_num}"
    tc_number = request.session[session_key]
    initial['tc_number'] = tc_number

    if request.method == 'POST':
        form_class = build_dynamic_form(template.id, data=request.POST)
        form = form_class(request.POST)

        if form.is_valid():
            tc_number = form.cleaned_data.get('tc_number', tc_number)
            if Document.objects.filter(document_number=tc_number).exists():
                form.add_error('tc_number', 'This TC number already exists. Please enter a unique number.')
                return render(request, 'documents/step3.html', {'form': form, 'template': template})

            serializable_data = {}
            for key, value in form.cleaned_data.items():
                if isinstance(value, (datetime.date, datetime.datetime)):
                    serializable_data[key] = value.isoformat()
                else:
                    serializable_data[key] = value

            # Create document (status='draft' initially)
            doc = Document.objects.create(
                client=client,
                product=product,
                template=template,
                document_number=tc_number,
                status='draft',
                created_by=request.user,
                form_data=serializable_data
            )

            # Save DocumentValue records
            for field_name, value in form.cleaned_data.items():
                try:
                    tfield = TemplateField.objects.get(template=template, field_name=field_name)
                except TemplateField.DoesNotExist:
                    continue
                if isinstance(value, (datetime.date, datetime.datetime)):
                    str_value = value.isoformat()
                else:
                    str_value = str(value)
                DocumentValue.objects.create(document=doc, field=tfield, value=str_value)

            # Clear the reserved TC number from session
            if session_key in request.session:
                del request.session[session_key]

            # ---- SYNCHRONOUS GENERATION (IMMEDIATE) ----
            success = _generate_document(doc)
            if not success:
                messages.error(request, 'Document generation failed. Please check logs.')
            else:
                messages.success(request, f'Document {doc.document_number} generated successfully.')

            return redirect('document_status', doc.id)
        else:
            return render(request, 'documents/step3.html', {'form': form, 'template': template})
    else:
        form_class = build_dynamic_form(template.id, initial=initial)
        form = form_class()
        return render(request, 'documents/step3.html', {'form': form, 'template': template})


# ---------- Document Status ----------
@login_required
def document_status(request, doc_id):
    doc = get_object_or_404(Document, id=doc_id)
    return render(request, 'documents/status.html', {'doc': doc})


# ---------- Document Preview ----------
@login_required
def document_preview(request, doc_id):
    doc = get_object_or_404(Document, id=doc_id)
    context = {val.field.field_name: val.value for val in doc.values.all()}
    context['company'] = doc.client.company
    context['document'] = doc
    html = DjangoTemplate(doc.template.html_layout).render(Context(context))
    return HttpResponse(html)


# ---------- Document List (with pagination & filters) ----------
@login_required
def document_list(request):
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    search = request.GET.get('search', '')
    status_filter = request.GET.get('status', '')

    docs = Document.objects.all().order_by('-created_at')

    if search:
        docs = docs.filter(
            Q(document_number__icontains=search) |
            Q(client__name__icontains=search) |
            Q(product__name__icontains=search) |
            Q(template__name__icontains=search)
        )
    if status_filter:
        docs = docs.filter(status=status_filter)
    if date_from:
        date_from = parse_date(date_from)
        if date_from:
            docs = docs.filter(created_at__date__gte=date_from)
    if date_to:
        date_to = parse_date(date_to)
        if date_to:
            docs = docs.filter(created_at__date__lte=date_to)

    paginator = Paginator(docs, 100)
    page = request.GET.get('page')
    try:
        docs_page = paginator.page(page)
    except PageNotAnInteger:
        docs_page = paginator.page(1)
    except EmptyPage:
        docs_page = paginator.page(paginator.num_pages)

    context = {
        'documents': docs_page,
        'search': search,
        'status_filter': status_filter,
        'date_from': date_from,
        'date_to': date_to,
    }
    return render(request, 'documents/list.html', context)


# ---------- Search ----------
@login_required
def search_documents(request):
    q = request.GET.get('q', '')
    docs = Document.objects.filter(
        Q(document_number__icontains=q) |
        Q(client__name__icontains=q) |
        Q(product__name__icontains=q) |
        Q(template__name__icontains=q)
    )
    return render(request, 'documents/search.html', {'documents': docs, 'query': q})


# ---------- Dashboard ----------
@login_required
def dashboard(request):
    today = timezone.now().date()
    month_start = today.replace(day=1)
    total_today = Document.objects.filter(created_at__date=today).count()
    total_month = Document.objects.filter(created_at__date__gte=month_start).count()
    pending = Document.objects.filter(status='draft').count()
    recent = Document.objects.order_by('-created_at')[:10]
    context = {
        'total_today': total_today,
        'total_month': total_month,
        'pending': pending,
        'recent': recent,
    }
    return render(request, 'dashboard.html', context)


# ---------- Public Verification Form ----------
def verify_form(request):
    if request.method == 'POST':
        doc_number = request.POST.get('document_number', '').strip()
        if doc_number:
            try:
                doc = Document.objects.get(document_number=doc_number, status='generated')
                return redirect('verify', document_number=doc_number)
            except Document.DoesNotExist:
                return render(request, 'verification/form.html', {
                    'error': 'Document not found or not yet generated.',
                    'document_number': doc_number
                })
        else:
            return render(request, 'verification/form.html', {'error': 'Please enter a document number.'})
    return render(request, 'verification/form.html')


# ---------- Public Verification Details ----------
def verify_document(request, document_number):
    doc = get_object_or_404(Document, document_number=document_number, status='generated')
    context = {val.field.field_name: val.value for val in doc.values.all()}
    context['document'] = doc
    return render(request, 'verification/verify.html', context)


# ---------- Download PDF ----------
@login_required
def download_pdf(request, doc_id):
    doc = get_object_or_404(Document, id=doc_id)
    if doc.generated_pdf:
        response = HttpResponse(doc.generated_pdf.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{doc.document_number}.pdf"'
        return response
    return HttpResponse('PDF not ready yet', status=404)


# ---------- Download DOCX ----------
@login_required
def download_docx(request, doc_id):
    doc = get_object_or_404(Document, id=doc_id)
    if doc.generated_docx:
        response = HttpResponse(doc.generated_docx.read(),
                                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = f'attachment; filename="{doc.document_number}.docx"'
        return response
    return HttpResponse('DOCX not ready yet', status=404)

def json_safe(obj):
    """Convert objects to JSON‑serializable format."""
    if isinstance(obj, datetime.datetime):
        return obj.date().isoformat()
    if isinstance(obj, datetime.date):
        return obj.isoformat()
    if isinstance(obj, datetime.time):
        return obj.isoformat()
    if isinstance(obj, dict):
        return {k: json_safe(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [json_safe(v) for v in obj]
    if isinstance(obj, str):
        if "T" in obj:
            try:
                return datetime.datetime.fromisoformat(obj).date().isoformat()
            except ValueError:
                pass
        return obj
    if obj is None:
        return ""
    return obj


@login_required
def upload_excel(request):
    if request.method == "POST":
        form = ExcelUploadForm(request.POST, request.FILES)

        if form.is_valid():
            excel_file = request.FILES["excel_file"]
            default_template = form.cleaned_data.get("template")

            try:
                wb = openpyxl.load_workbook(excel_file)
            except Exception as e:
                messages.error(request, f"Invalid Excel file: {e}")
                return redirect("upload_excel")

            sheet = wb.active
            headers = [cell.value for cell in sheet[1]]

            if not headers:
                messages.error(request, "The Excel file is empty.")
                return redirect("upload_excel")

            today = datetime.date.today().isoformat()
            success_count = 0
            errors = []

            for row_idx, row in enumerate(
                sheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                if not any(row):
                    continue

                row_data = dict(zip(headers, row))
                row_data = json_safe(row_data)

                client_name = row_data.get("client_name")
                product_name = row_data.get("product_name")
                template_name = row_data.get("template_name")
                tc_number = row_data.get("tc_number")

                if not row_data.get("date"):
                    row_data["date"] = today
                if not row_data.get("testing_date"):
                    row_data["testing_date"] = today

                if not client_name:
                    errors.append(f"Row {row_idx}: client_name is required.")
                    continue
                if not product_name:
                    errors.append(f"Row {row_idx}: product_name is required.")
                    continue

                # ----- Client -----
                try:
                    client = Client.objects.get(name__iexact=client_name, is_active=True)
                except Client.DoesNotExist:
                    company = Company.objects.first()
                    if not company:
                        errors.append(f'Row {row_idx}: No company found for client "{client_name}".')
                        continue
                    client = Client.objects.create(
                        company=company,
                        name=client_name,
                        address="",
                        email="",
                        phone="",
                        contact_person="",
                        is_active=True,
                    )
                    messages.info(request, f'Created client "{client_name}".')

                # ----- Product -----
                try:
                    product = Product.objects.get(name__iexact=product_name)
                except Product.DoesNotExist:
                    errors.append(f'Row {row_idx}: Product "{product_name}" not found.')
                    continue

                # ----- Template -----
                if template_name:
                    try:
                        template = Template.objects.get(name__iexact=template_name, is_active=True)
                    except Template.DoesNotExist:
                        errors.append(f'Row {row_idx}: Template "{template_name}" not found.')
                        continue
                elif default_template:
                    template = default_template
                else:
                    template = product.default_template

                if not template:
                    errors.append(f"Row {row_idx}: No template selected.")
                    continue

                # ----- TC Number Assignment -----
                if not tc_number:
                    prefix = "AFPL/2026-27/"
                    tc_number = f"{prefix}{get_next_number(prefix)}"
                else:
                    if Document.objects.filter(document_number=tc_number).exists():
                        errors.append(f'Row {row_idx}: TC Number "{tc_number}" already exists.')
                        continue

                # ----- Build form_data from all columns -----
                form_data = {}
                for header in headers:
                    if header == "template_name":
                        continue
                    value = row_data.get(header)
                    if value is None:
                        value = ""
                    form_data[header] = json_safe(value)

                # 🔥 IMPORTANT: Ensure tc_number is in form_data
                form_data["tc_number"] = tc_number
                form_data["date"] = json_safe(form_data.get("date") or today)
                form_data["testing_date"] = json_safe(form_data.get("testing_date") or today)

                # Verify JSON safety
                try:
                    json.dumps(form_data)
                except Exception as e:
                    errors.append(f"Row {row_idx}: JSON Error -> {e}")
                    continue

                # ----- Save Document -----
                try:
                    with transaction.atomic():
                        document = Document.objects.create(
                            client=client,
                            product=product,
                            template=template,
                            document_number=tc_number,
                            status="draft",
                            created_by=request.user,
                            form_data=form_data,
                        )

                        template_fields = {
                            f.field_name: f
                            for f in TemplateField.objects.filter(template=template)
                        }

                        values = []
                        for field_name, value in form_data.items():
                            field = template_fields.get(field_name)
                            if not field:
                                continue
                            values.append(
                                DocumentValue(
                                    document=document,
                                    field=field,
                                    value=str(value),
                                )
                            )

                        if values:
                            DocumentValue.objects.bulk_create(values)

                        # Generate PDF/DOCX synchronously
                        success = _generate_document(document)
                        if success:
                            success_count += 1
                        else:
                            errors.append(f"Row {row_idx}: Failed to generate document.")

                except Exception as e:
                    errors.append(f"Row {row_idx}: {e}")

            # Show results
            for err in errors:
                messages.error(request, err)
            if success_count:
                messages.success(request, f"Successfully processed {success_count} rows.")
            return redirect("document_list")

    else:
        form = ExcelUploadForm()

    return render(request, "documents/upload_excel.html", {"form": form})


# ---------- Download Sample Excel ----------
@login_required
def download_sample_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Documents"

    headers = [
        'client_name', 'product_name', 'template_name', 'tc_number',
        'client_po_no', 'po_date', 'filter_serial_number', 'testing_date',
        'equipment_used', 'filter_type', 'filter_media', 'filter_capacity',
        'filter_dimension', 'moc', 'air_flow_rate', 'initial_pressure_drop',
        'final_pressure_drop', 'temperature', 'test_airflow', 'air_velocity',
        'air_temperature', 'test_initial_pressure_drop', 'efficiency',
        'filter_class', 'test_type'
    ]
    ws.append(headers)

    sample = [
        'Samarth Engineering Services', 'Fine Filter',
        'Test Certificate', '',
        'PO-12345', '2026-07-14', 'SER/001', '2026-07-14',
        'VANE TYPE ANEMOMETER AND AEROSOL PHOTOMETER',
        'FINE FILTER -- 3 MICRON (FLANGE TYPE)',
        'HDPE + NON-WOVEN SYNTHETIC MEDIA + HDPE',
        '2000 CFM', '610 X 610 X 305 MM', 'ALUMINIUM',
        '2000 CFM', '70 - 80 Pa', '450 Pa', 'AMBIENT',
        '2000 CFM ± 5%', '500 FPM', '25 °C',
        '70 ± 5 %', '99% DOWN TO 3 MICRON', 'EU 07',
        'BS EN 779'
    ]
    ws.append(sample)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=sample_documents.xlsx'
    response.write(save_virtual_workbook(wb))
    return response


# ---------- Regenerate Single Document ----------
@login_required
def regenerate_document(request, doc_id):
    doc = get_object_or_404(Document, id=doc_id)

    if doc.status not in ['generated', 'failed', 'draft']:
        messages.warning(request, 'Document is not in a valid state for regeneration.')
        return redirect('document_status', doc.id)

    if doc.generated_pdf:
        doc.generated_pdf.delete(save=False)
    if doc.generated_docx:
        doc.generated_docx.delete(save=False)
    if doc.qr_code:
        doc.qr_code.delete(save=False)
    if doc.barcode:
        doc.barcode.delete(save=False)

    doc.status = 'draft'
    doc.save()

    success = _generate_document(doc)
    if success:
        messages.success(request, f'Document {doc.document_number} regenerated successfully.')
    else:
        messages.error(request, f'Failed to regenerate document {doc.document_number}.')

    return redirect('document_status', doc.id)


# ---------- Bulk Actions ----------
@login_required
@require_POST
def bulk_action(request):
    action = request.POST.get('action')
    doc_ids = request.POST.getlist('doc_ids')
    if not doc_ids:
        messages.error(request, 'No documents selected.')
        return redirect('document_list')

    if action == 'download_pdf':
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w') as zip_file:
            for doc_id in doc_ids:
                doc = get_object_or_404(Document, id=doc_id)
                if doc.generated_pdf and doc.status == 'generated':
                    pdf_content = doc.generated_pdf.read()
                    zip_file.writestr(f"{doc.document_number}.pdf", pdf_content)
        zip_buffer.seek(0)
        response = HttpResponse(zip_buffer, content_type='application/zip')
        response['Content-Disposition'] = 'attachment; filename="documents_pdf.zip"'
        return response

    elif action == 'download_docx':
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w') as zip_file:
            for doc_id in doc_ids:
                doc = get_object_or_404(Document, id=doc_id)
                if doc.generated_docx and doc.status == 'generated':
                    docx_content = doc.generated_docx.read()
                    zip_file.writestr(f"{doc.document_number}.docx", docx_content)
        zip_buffer.seek(0)
        response = HttpResponse(zip_buffer, content_type='application/zip')
        response['Content-Disposition'] = 'attachment; filename="documents_docx.zip"'
        return response

    # === NEW: Merge PDFs into a single file ===
    elif action == 'merge_pdfs':
        from PyPDF2 import PdfMerger
        merger = PdfMerger()
        merged_filenames = []   # will store full document numbers
        number_map = {}         # maps numeric suffix -> full doc number

        for doc_id in doc_ids:
            doc = get_object_or_404(Document, id=doc_id)
            if doc.status == 'generated' and doc.generated_pdf:
                pdf_bytes = doc.generated_pdf.read()
                merger.append(io.BytesIO(pdf_bytes))
                merged_filenames.append(doc.document_number)
                # Extract numeric suffix (last part after '/')
                try:
                    num_part = doc.document_number.split('/')[-1]
                    if num_part.isdigit():
                        number_map[int(num_part)] = doc.document_number
                except:
                    pass

        if not merged_filenames:
            messages.error(request, 'No valid generated PDFs found to merge.')
            return redirect('document_list')

        # ---- Generate a meaningful filename ----
        if number_map:
            sorted_numbers = sorted(number_map.keys())
            base_prefix = merged_filenames[0].split('/')[0] + '_' + merged_filenames[0].split('/')[1]  # e.g., "AFPL_2026-27"
            # if only one
            if len(sorted_numbers) == 1:
                filename = f"{base_prefix}_{sorted_numbers[0]}.pdf"
            else:
                # Check if numbers are consecutive
                consecutive = all(sorted_numbers[i+1] - sorted_numbers[i] == 1 for i in range(len(sorted_numbers)-1))
                if consecutive:
                    filename = f"{base_prefix}_{sorted_numbers[0]}_to_{sorted_numbers[-1]}.pdf"
                else:
                    # join all numbers with underscores
                    num_str = '_'.join(str(n) for n in sorted_numbers)
                    filename = f"{base_prefix}_{num_str}.pdf"
        else:
            # fallback
            filename = "merged_documents.pdf"

        # Replace slashes with underscores to avoid filesystem issues
        filename = filename.replace('/', '_')

        # ---- Return the merged PDF ----
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        merger.write(response)
        merger.close()
        return response

    elif action == 'regenerate':
        count = 0
        for doc_id in doc_ids:
            doc = get_object_or_404(Document, id=doc_id)
            if doc.status == 'generated':
                doc.status = 'draft'
                doc.generated_pdf.delete(save=False)
                doc.generated_docx.delete(save=False)
                doc.qr_code.delete(save=False)
                doc.barcode.delete(save=False)
                doc.save()
                generate_document_pdf_docx(doc.id)
                count += 1
        messages.success(request, f'Regeneration started for {count} documents.')
        return redirect('document_list')

    else:
        messages.error(request, 'Invalid action.')
        return redirect('document_list')


@csrf_exempt
def process_queue(request):
    """
    Process pending documents synchronously – no background worker required.
    Called by external cron job (e.g., cron-job.org).
    """
    MAX_PER_RUN = 25
    processed = 0
    errors = 0

    pending_docs = Document.objects.filter(
        Q(status='draft') | Q(status='processing')
    ).order_by('created_at')[:MAX_PER_RUN]

    for doc in pending_docs:
        if doc.status == 'processing':
            doc.status = 'draft'
            doc.save()

        try:
            success = _generate_document(doc)
            if success:
                processed += 1
            else:
                errors += 1
        except Exception as e:
            logger.error(f"Error processing doc {doc.id}: {e}")
            doc.status = 'failed'
            doc.save()
            errors += 1

    remaining = Document.objects.filter(status='draft').count()
    return JsonResponse({
        'status': 'ok',
        'processed': processed,
        'errors': errors,
        'remaining': remaining
    })


@login_required
def document_edit(request, doc_id):
    doc = get_object_or_404(Document, id=doc_id)
    template = doc.template

    initial = {}
    for val in doc.values.all():
        initial[val.field.field_name] = val.value

    if request.method == 'POST':
        form_class = build_dynamic_form(template.id, data=request.POST)
        form = form_class(request.POST)
        if form.is_valid():
            doc.values.all().delete()
            for field_name, value in form.cleaned_data.items():
                try:
                    tfield = TemplateField.objects.get(template=template, field_name=field_name)
                except TemplateField.DoesNotExist:
                    continue
                if isinstance(value, (datetime.date, datetime.datetime)):
                    str_value = value.isoformat()
                else:
                    str_value = str(value)
                DocumentValue.objects.create(document=doc, field=tfield, value=str_value)

            doc.status = 'draft'
            doc.generated_pdf.delete(save=False)
            doc.generated_docx.delete(save=False)
            doc.qr_code.delete(save=False)
            doc.barcode.delete(save=False)
            doc.save()

            generate_document_pdf_docx(doc.id)
            messages.success(request, 'Document updated and regeneration started.')
            return redirect('document_status', doc.id)
        else:
            return render(request, 'documents/edit.html', {'form': form, 'document': doc})
    else:
        form_class = build_dynamic_form(template.id, initial=initial)
        form = form_class()
        return render(request, 'documents/edit.html', {'form': form, 'document': doc})


